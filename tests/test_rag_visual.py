import io,json,tempfile,threading,unittest
from pathlib import Path
from unittest.mock import patch
from h3chat.service import Service
from h3chat.store import DEFAULTS
from h3chat.rag_visual import attach
from h3chat.ovis import messages

ROOT=Path(__file__).resolve().parents[1]

class VisualRagTests(unittest.TestCase):
    def setUp(self):
        import sys
        sys.path.insert(0,str(ROOT/'runtime/tools/documents'))
        self.tmp=tempfile.TemporaryDirectory();self.app=Service(ROOT,self.tmp.name,start_worker=False)
        self.addCleanup(self.tmp.cleanup);self.addCleanup(self.app.close)
        self.project=self.app.knowledge.save({'name':'Visual fixture'});self.cancel=threading.Event()
        self.settings=DEFAULTS|{'rag_embedding_profile':'ovis','rag_embedding_model':'synthetic'}

    def picture(self):
        from PIL import Image,ImageDraw
        image=Image.new('RGB',(200,160),'white');ImageDraw.Draw(image).ellipse((25,20,175,140),fill='red')
        stream=io.BytesIO();image.save(stream,format='PNG');return stream.getvalue()

    def link(self,path):
        with patch.object(self.app.knowledge,'refresh'):self.app.knowledge.add(self.project['id'],{'paths':[str(path)]})
        return self.app.knowledge.project(self.project['id'])['sources'][-1]

    def test_real_pdf_photo_scan_and_vector_diagram_keep_page_provenance(self):
        from PIL import Image
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject,NameObject,NumberObject,DecodedStreamObject
        writer=PdfWriter();page=writer.add_blank_page(width=300,height=400)
        image=DecodedStreamObject();image.set_data(Image.open(io.BytesIO(self.picture())).convert('RGB').tobytes())
        image.update({NameObject('/Type'):NameObject('/XObject'),NameObject('/Subtype'):NameObject('/Image'),NameObject('/Width'):NumberObject(200),NameObject('/Height'):NumberObject(160),NameObject('/ColorSpace'):NameObject('/DeviceRGB'),NameObject('/BitsPerComponent'):NumberObject(8)})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/XObject'):DictionaryObject({NameObject('/Im1'):writer._add_object(image)})})
        content=DecodedStreamObject();content.set_data(b'q 200 0 0 160 20 100 cm /Im1 Do Q');page[NameObject('/Contents')]=writer._add_object(content)
        page=writer.add_blank_page(width=300,height=400);content=DecodedStreamObject();content.set_data(b'0 0 1 rg 40 100 180 180 re f');page[NameObject('/Contents')]=writer._add_object(content)
        writer.add_blank_page(width=300,height=400)
        path=Path(self.tmp.name)/'illustrated.pdf'
        with path.open('wb') as out:writer.write(out)
        source=self.link(path);chunks,warnings=self.app.knowledge.read_visuals(source,[],self.cancel,lambda _:None)
        self.assertEqual([b['page'] for b in chunks],[1,2]);self.assertTrue(all(b['modality']=='image' for b in chunks))
        self.assertTrue(all((self.app.data/b['image_path']).is_file() for b in chunks));self.assertEqual(warnings,[])
        # A scan with no extractable text still becomes searchable visually.
        with patch.object(self.app.knowledge.embeddings,'key',return_value='visual-fixture'),patch.object(self.app.knowledge.embeddings,'encode_items',side_effect=lambda items,*a:[[1.]+[0.]*7 for _ in items]):
            self.app.knowledge.index(self.project['id'],self.settings,self.cancel,lambda _:None)
        self.assertEqual(self.app.knowledge.project(self.project['id'])['sources'][0]['status'],'ready')
        self.assertEqual(self.app.store.one('SELECT COUNT(*) AS n FROM rag_chunks')['n'],2)

    def test_word_images_and_text_preserve_block_location(self):
        from docx import Document
        doc=Document();doc.add_paragraph('A verified caption');doc.add_picture(io.BytesIO(self.picture()))
        path=Path(self.tmp.name)/'illustrated.docx';doc.save(path);source=self.link(path)
        chunks,_=self.app.knowledge.read_visuals(source,[{'location':'blocco 2','text':'Caption next to image','page':None}],self.cancel,lambda _:None)
        self.assertEqual(len(chunks),1);self.assertEqual(chunks[0]['location'],'blocco 2 · immagine');self.assertIn('Caption next to image',chunks[0]['text'])

    def test_semantic_image_retrieval_and_cached_index(self):
        path=Path(self.tmp.name)/'circle.png';path.write_bytes(self.picture());self.link(path)
        calls=[]
        def encode(items,*args):
            calls.extend(items);return [[1.]+[0.]*7 for _ in items]
        with patch.object(self.app.knowledge.embeddings,'key',return_value='fixture'),patch.object(self.app.knowledge.embeddings,'encode_items',side_effect=encode),patch.object(self.app.knowledge.embeddings,'encode',return_value=[[1.]+[0.]*7]):
            rows,_=self.app.knowledge.retrieve(self.project['id'],'cerchio rosso',self.settings,self.cancel,lambda _:None)
            self.assertEqual(rows[0]['modality'],'image');self.assertTrue(rows[0]['image_path']);self.assertEqual(len(calls),1)
            self.app.knowledge.retrieve(self.project['id'],'red circle',self.settings,self.cancel,lambda _:None);self.assertEqual(len(calls),1)
            (self.app.data/rows[0]['image_path']).unlink()
            repaired,_=self.app.knowledge.retrieve(self.project['id'],'red circle',self.settings,self.cancel,lambda _:None)
            self.assertTrue((self.app.data/repaired[0]['image_path']).is_file());self.assertEqual(len(calls),2)
        self.assertTrue(path.exists())
        other=self.app.knowledge.save({'name':'Other'})
        with patch.object(self.app.knowledge.embeddings,'key',return_value='fixture'),patch.object(self.app.knowledge.embeddings,'encode',return_value=[[1.]+[0.]*7]):
            self.assertEqual(self.app.knowledge.retrieve(other['id'],'circle',self.settings,self.cancel,lambda _:None)[0],[])

    def test_vision_handoff_is_durable_bounded_and_explicit_when_disabled(self):
        original=self.app.data/'project-cache/circle.png';original.parent.mkdir();original.write_bytes(self.picture())
        self.app.catalog['vision-fixture']={'max_refs':1,'vision':{'enabled':True}}
        sources=[{'citation':'R'+str(i),'name':'Book','location':f'pagina {i}','image_path':'project-cache/circle.png'} for i in (1,2)]
        history=[{'content':'Explain','media':[]}]
        result=attach(self.app,{'id':'fixture'},history,sources,DEFAULTS|{'chat_model':'vision-fixture'})
        self.assertEqual(len(history[-1]['media']),1);self.assertEqual(result['rag_visual_count'],1);self.assertIn('1 immagine',result['rag_visual_warning'])
        self.assertIn('[R1]',history[-1]['content']);self.assertTrue(all((self.app.data/s['image']['path']).is_file() for s in sources))
        original.unlink();self.assertTrue((self.app.data/sources[0]['image']['path']).is_file())

    def test_nonvision_model_gets_no_claim_of_visual_analysis(self):
        original=self.app.data/'project-cache/circle.png';original.parent.mkdir();original.write_bytes(self.picture())
        sources=[{'citation':'R1','name':'Book','location':'pagina 1','image_path':'project-cache/circle.png'}];history=[{'content':'Explain','media':[]}]
        result=attach(self.app,{'id':'fixture'},history,sources,DEFAULTS|{'vision_enabled':False})
        self.assertEqual(result['rag_visual_count'],0);self.assertEqual(history[-1]['media'],[]);self.assertIn('Non inventare',history[-1]['content'])

    def test_quick_device_options_persist_and_refuse_active_jobs(self):
        self.app.save_rag_options({'rag_device':'gpu','rag_visual':False})
        self.assertEqual(self.app.store.settings()['rag_device'],'gpu');self.assertFalse(self.app.store.settings()['rag_visual'])
        for patch_value in ({'rag_device':'bad'},{'rag_visual':'yes'},{'chat_model':'other'}):
            with self.assertRaises(ValueError):self.app.save_rag_options(patch_value)
        chat=self.app.store.create_chat();self.app.store.enqueue(chat['id'],'Synthetic',[],DEFAULTS)
        with self.assertRaisesRegex(ValueError,'lavori'):self.app.save_rag_options({'rag_device':'cpu'})

    def test_background_gpu_index_serializes_generation_and_wait_is_cancellable(self):
        import time
        self.app.store.save_settings({'rag_device':'gpu','rag_embedding_model':'synthetic'})
        entered=threading.Event();release=threading.Event()
        def slow(*args,**kw):entered.set();release.wait(3)
        with patch.object(self.app.knowledge,'index',side_effect=slow),patch.object(self.app.knowledge,'release_gpu') as stop:
            self.app.knowledge.refresh(self.project['id']);self.assertTrue(entered.wait(2));self.assertTrue(stop.called)
            with self.assertRaisesRegex(ValueError,'completamento'):self.app.save_rag_options({'rag_device':'cpu'})
            chat=self.app.store.create_chat();ident=self.app.store.enqueue(chat['id'],'Synthetic',[],DEFAULTS)
            job=self.app.store.one('SELECT * FROM jobs WHERE id=?',(ident,));cancel=threading.Event()
            thread=threading.Thread(target=self.app.execute_job,args=(job,cancel));thread.start();time.sleep(.25);cancel.set();thread.join(1)
            self.assertFalse(thread.is_alive());self.assertEqual(self.app.store.one('SELECT status FROM jobs WHERE id=?',(ident,))['status'],'cancelled')
            release.set()
            for _ in range(100):
                if not self.app.knowledge.project(self.project['id'])['indexing']:break
                time.sleep(.01)
            self.assertFalse(self.app.knowledge.project(self.project['id'])['indexing'])

    def test_visual_prompt_treats_filename_and_caption_as_data(self):
        value=messages('Synthetic caption',False,'F:/book/image.png')
        self.assertEqual(value[1]['content'][0],{'type':'image','image':'F:/book/image.png'})
        self.assertEqual(value[1]['content'][1]['text'],'Synthetic caption')

    def test_chat_job_passes_retrieved_image_to_vision_and_keeps_citation_preview(self):
        path=Path(self.tmp.name)/'circle.png';path.write_bytes(self.picture());source=self.link(path)
        original=self.app.data/'project-cache/circle.png';original.parent.mkdir();original.write_bytes(self.picture())
        model={'id':'vision-fixture','name':'Vision fixture','capabilities':['chat'],'files':[],'vision':{'enabled':True,'max_refs':1}}
        self.app.catalog[model['id']]=model;self.app.store.save_settings({'chat_model':model['id']})
        chat=self.app.store.create_chat(project_id=self.project['id'])
        row={'id':1,'source_id':source['id'],'name':'Illustrated book','location':'pagina 2 · contenuto visivo','page':2,'text':'Contenuto visivo del documento.','source_url':'','image_path':'project-cache/circle.png'}
        with patch.object(self.app.engine,'require_model',return_value=model):ticket=self.app.send(chat['id'],{'prompt':'Descrivi il cerchio','rag':True})
        self.app.catalog[model['id']]=model
        job=self.app.store.one('SELECT * FROM jobs WHERE id=?',(ticket['job_id'],))
        with patch.object(self.app.knowledge,'retrieve',return_value=([row],'ibrida · GPU')),patch.object(self.app.engine,'require_model',return_value=model),patch.object(self.app.engine,'start_llama'),patch.object(self.app.engine,'completion',return_value=('Un cerchio rosso [R1].','stop')) as completion:
            self.app.execute_job(job,self.cancel)
        answer=self.app.store.messages(chat['id'])[-1];self.assertEqual(answer['status'],'done')
        self.assertEqual(answer['meta']['rag_visual_count'],1);self.assertEqual(answer['meta']['rag_sources'][0]['page'],2)
        self.assertEqual(len([x for x in completion.call_args.args[0][-1]['content'] if x['type']=='image_url']),1)
        self.assertTrue((self.app.data/answer['meta']['rag_sources'][0]['image']['path']).is_file())

if __name__=='__main__':unittest.main()
