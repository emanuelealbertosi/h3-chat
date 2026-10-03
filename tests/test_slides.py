import base64
import io
import json
import sys
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import patch

from h3chat.service import Service
from h3chat.slides import PREFIX, requested, options, validate_page, normalize_page, partial_page, validate_content, encode,edit_request,edit_options
from h3chat.pdf_export import export_html

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'runtime/tools/documents'))

def node(kind='text',text='Fonte [R1]',**kw):
    return {'id':'n1','parent':'root','kind':kind,'text':text,'asset_id':'','language':'text','style':{'flow':'stack','columns':[1,1],'surface':'plain','gap':20}}|kw

class SchemaTests(unittest.TestCase):
    def test_route_and_options(self):
        for prompt in ('Crea 12 slide dal PDF','Trasforma queste immagini in una presentazione','Prepara le diapositive dal RAG','Modifica la presentazione in 4:3'):
            self.assertTrue(requested(prompt))
        for prompt in ('Come si creano le slide?','Analizza queste slide','Crea un’immagine del PowerPoint aperto'):
            # The last request explicitly describes an image; media selection
            # takes precedence at the service boundary.
            if 'immagine' not in prompt:self.assertFalse(requested(prompt))
        self.assertEqual(options('Crea 12 slide in 4:3'),{'count':12,'format':'4:3'})
        for count in (0,31,True,'8'):
            with self.assertRaises(ValueError):options('',{'count':count})

    def test_incremental_text_and_final_validation(self):
        raw='{"nodes":[{"id":"n1","parent":"root","kind":"text","text":"A metà '
        self.assertEqual(partial_page(raw,set(),set())['nodes'][0]['text'],'A metà ')
        raw=json.dumps({'nodes':[node(text='Una formula $x^2$')],'notes':'','sources':['R1']})
        self.assertEqual(validate_page(json.loads(raw),set(),{'R1'})['sources'],['R1'])
        for n in (node(id='root'),node(parent='missing'),node(asset_id='untrusted',kind='image'),node(style={'flow':'javascript'})):
            with self.assertRaises(ValueError):validate_page({'nodes':[n]},set(),set())
        with self.assertRaises(ValueError):validate_page({'nodes':[node()], 'sources':['inventata']},set(),{'R1'})
        escaped='{"nodes":[{"id":"n1","parent":"root","kind":"text","text":"He said \\"hello\\"'
        self.assertIn('He said',partial_page(escaped,set(),set())['nodes'][0]['text'])

    def test_editing_preserves_format_and_resolves_addition(self):
        content=encode({'version':1,'format':'4:3','pages':[{'title':'A','nodes':[]},{'title':'B','nodes':[]}]})
        self.assertTrue(edit_request('Riduci il testo della seconda pagina',content))
        self.assertEqual(edit_options('Riduci il testo',content),{'count':2,'format':'4:3'})
        self.assertEqual(edit_options('Aggiungi 2 slide in 16:9',content),{'count':4,'format':'16:9'})

    def test_model_labels_forward_groups_and_root_container(self):
        raw={'nodes':[node(id='2. Testo',parent='layout principale'),
                      node('group',id='layout principale',parent='root'),
                      node('group',id='root',parent=None,style={'flow':'columns','columns':[2,1]})]}
        page=normalize_page(raw,set(),set())
        self.assertEqual([n['id'] for n in page['nodes']],['n3','n2','n1'])
        self.assertEqual([n['parent'] for n in page['nodes']],['root','n3','n2'])
        self.assertEqual(page['nodes'][0]['style']['flow'],'columns')
        validate_page(page,set(),set())
        for nodes in ([node(id='duplicate'),node(id='duplicate')],
                      [node(parent='not provided')],
                      [node('group',parent='n2'),node('group',id='n2',parent='n1')],
                      [node(id='root')]):
            with self.assertRaises(ValueError):normalize_page({'nodes':nodes},set(),set())

    def test_preview_reads_reordered_properties_and_unfinished_groups(self):
        raw='{"nodes":[{"kind":"text","text":"Contenuto in tempo reale'
        self.assertEqual(partial_page(raw,set(),set())['nodes'][0]['text'],'Contenuto in tempo reale')
        raw='{"nodes":[{"text":"Testo prima del tipo'
        self.assertEqual(partial_page(raw,set(),set())['nodes'][0]['text'],'Testo prima del tipo')
        raw='{"nodes":['+json.dumps(node(id='1. Titolo',parent='gruppo futuro'))+',{"kind":"text","text":"Secondo '
        draft=partial_page(raw,set(),set())
        self.assertEqual([n['text'] for n in draft['nodes']],['Fonte [R1]','Secondo '])
        self.assertTrue(all(n['parent']=='root' for n in draft['nodes']))
        # Words resembling fields inside a string must not become node fields.
        raw='{"nodes":[{"kind":"text","text":"Una chiave \\"parent\\": \\"evil\\"'
        self.assertEqual(partial_page(raw,set(),set())['nodes'][0]['parent'],'root')

    def test_every_token_of_noncanonical_page_can_be_previewed(self):
        raw=json.dumps({'nodes':[node('group',id='0.layout'),node(id='1.text',parent='0.layout',text='Testo progressivo con $x^2$ e "virgolette".')],'notes':'','sources':[]})
        longest=0
        for size in range(1,len(raw)+1):
            draft=partial_page(raw[:size],set(),set())
            if draft:
                validate_page(draft,set(),set(),draft=True)
                longest=max(longest,sum(len(n['text']) for n in draft['nodes']))
        self.assertGreaterEqual(longest,len('Testo progressivo con $x^2$ e "virgolette".'))

class GenerationTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.app=Service(ROOT,Path(self.temp.name),start_worker=False);self.addCleanup(self.app.close)
        self.model={'id':'fixture','name':'Current LLM','files':[],'capabilities':['chat'],'vision':{'enabled':False}}
        self.cancel=threading.Event()

    def enqueue(self,prompt='Crea 2 slide dal documento',**body):
        chat=self.app.store.create_chat()
        with patch.object(self.app.engine,'require_model',return_value=self.model):
            sent=self.app.send(chat['id'],{'prompt':prompt,**body})
        return self.app.store.one('SELECT * FROM jobs WHERE id=?',(sent['job_id'],))

    def completion(self,messages,settings,cancel,**kw):
        self.calls.append(messages)
        schema=kw['schema']
        if 'slides' in schema['properties']:
            value={'title':'Report','slides':[{'title':'Pagina '+str(i),'purpose':'Dati del documento'} for i in range(2)]}
        else:
            draft='{"nodes":[{"id":"n1","parent":"root","kind":"text","text":"Progressivo'
            kw['on_text'](draft)
            saved=self.app.store.canvas_history.get(self.job['chat_id'])
            self.assertIn('Progressivo',saved['content']);self.progress+=1
            if self.fail:raise RuntimeError('Synthetic interruption')
            value={'nodes':[node(text='Entrate 1200 dal documento')],'notes':'Fonte fornita','sources':['D1'] if self.documents else []}
        return json.dumps(value),'stop'

    def execute(self,job,*,fail=False,documents=False):
        self.calls=[];self.job=job;self.progress=0;self.fail=fail;self.documents=documents
        with patch.object(self.app.engine,'require_model',return_value=self.model),patch.object(self.app.engine,'prepare'),patch.object(self.app.engine,'start_llama') as start,patch.object(self.app.engine,'completion',side_effect=self.completion):
            self.app.execute_job(job,self.cancel)
        self.assertEqual(start.call_count,1)
        return self.app.store.messages(job['chat_id'])[-1]

    def test_incremental_canvas_history_and_no_body_artifact(self):
        job=self.enqueue();self.assertTrue(json.loads(job['payload'])['canvas'])
        answer=self.execute(job);self.assertEqual(answer['status'],'done');self.assertEqual(self.progress,2)
        self.assertEqual(answer['media'],[]);self.assertNotIn(PREFIX,answer['content'])
        self.assertEqual(len(self.app.store.canvas_history.listing(job['chat_id'])['items']),1)
        value=self.app.store.canvas_history.get(job['chat_id']);validate_content(value['content'],value['media'])
        deck=json.loads(value['content'][len(PREFIX):-4]);self.assertEqual(len(deck['pages']),2)
        self.assertTrue(all(p['status']=='ready' for p in deck['pages']))

    def test_failure_preserves_streaming_artifact(self):
        job=self.enqueue();answer=self.execute(job,fail=True)
        self.assertEqual(answer['status'],'failed');self.assertIn('Progressivo',answer['meta']['artifact']['content'])
        self.assertEqual(len(self.app.store.canvas_history.listing(job['chat_id'])['items']),1)

    def test_generation_recovers_invalid_tree_and_streams_before_node_id(self):
        job=self.enqueue();page_calls=0;self.calls=[];self.job=job;self.progress=0;self.fail=False;self.documents=False
        def complete(messages,settings,cancel,**kw):
            nonlocal page_calls
            if 'slides' in kw['schema']['properties']:return self.completion(messages,settings,cancel,**kw)
            page_calls+=1
            if page_calls==2:
                kw['on_text']('{"nodes":[{"id":"layout","parent":"root","kind":"group","text":""')
                self.assertIn('Già visibile',self.app.store.canvas_history.get(job['chat_id'])['content'])
            kw['on_text']('{"nodes":[{"kind":"text","text":"Già visibile')
            self.assertIn('Già visibile',self.app.store.canvas_history.get(job['chat_id'])['content'])
            if page_calls==1:return json.dumps({'nodes':[node(parent='unprovided')]}),'stop'
            if page_calls==2:self.assertIn('Correggi soltanto',messages[-1]['content'])
            return json.dumps({'nodes':[node(id='1. Testo',parent='contenitore'),node('group',id='contenitore')]}),'stop'
        with patch.object(self.app.engine,'require_model',return_value=self.model),patch.object(self.app.engine,'prepare'),patch.object(self.app.engine,'start_llama'),patch.object(self.app.engine,'completion',side_effect=complete):
            self.app.execute_job(job,self.cancel)
        answer=self.app.store.messages(job['chat_id'])[-1]
        self.assertEqual(answer['status'],'done',answer['meta'].get('error'));self.assertEqual(page_calls,3)
        validate_content(answer['meta']['artifact']['content'],[])

    def test_broken_tree_retry_is_bounded_and_retains_preview(self):
        job=self.enqueue();self.calls=[];self.job=job;self.progress=0;self.fail=False;self.documents=False;page_calls=0
        def complete(messages,settings,cancel,**kw):
            nonlocal page_calls
            if 'slides' in kw['schema']['properties']:return self.completion(messages,settings,cancel,**kw)
            page_calls+=1;kw['on_text']('{"nodes":[{"kind":"text","text":"Anteprima conservata')
            return json.dumps({'nodes':[node(parent='unknown')]}),'stop'
        with patch.object(self.app.engine,'require_model',return_value=self.model),patch.object(self.app.engine,'prepare'),patch.object(self.app.engine,'start_llama'),patch.object(self.app.engine,'completion',side_effect=complete):
            self.app.execute_job(job,self.cancel)
        answer=self.app.store.messages(job['chat_id'])[-1]
        self.assertEqual(answer['status'],'failed');self.assertEqual(page_calls,2)
        self.assertIn('Anteprima conservata',answer['meta']['artifact']['content'])

    def test_real_word_text_and_figures_reach_current_llm(self):
        from docx import Document
        from PIL import Image
        doc=Document();doc.add_paragraph('Entrate 1200. Utile 900.')
        image=io.BytesIO();Image.new('RGB',(160,100),'green').save(image,'PNG');image.seek(0);doc.add_picture(image)
        stream=io.BytesIO();doc.save(stream)
        upload=self.app.upload({'name':'Report.docx','data':base64.b64encode(stream.getvalue()).decode()})
        job=self.enqueue(media=[upload]);answer=self.execute(job,documents=True)
        self.assertEqual(answer['status'],'done',answer['meta'].get('error'))
        self.assertIn('Entrate 1200',json.dumps(self.calls,ensure_ascii=False))
        asset=answer['meta']['artifact']['media'][0];self.assertTrue(asset['path'].startswith('outputs/'));self.assertTrue((self.app.data/asset['path']).exists())
        self.assertIn('Report.docx',answer['meta']['artifact']['content'])

    def test_project_overview_uses_only_selected_rag_sources(self):
        project=self.app.knowledge.save({'name':'Fixture','instructions':'Usa euro.','sources_only':True})
        for ident,text in (('chosen','Valore verificato 1234 euro.'),('excluded','Dato escluso 999999.')):
            path=Path(self.temp.name)/(ident+'.md');path.write_text(text)
            self.app.store.execute('INSERT INTO project_sources(id,project_id,path,name) VALUES (?,?,?,?)',(ident,project['id'],str(path),path.name))
        chat=self.app.store.create_chat(project_id=project['id'])
        with patch.object(self.app.engine,'require_model',return_value=self.model):
            sent=self.app.send(chat['id'],{'prompt':'Crea 2 slide','rag':True,'rag_sources':['chosen']})
        job=self.app.store.one('SELECT * FROM jobs WHERE id=?',(sent['job_id'],));answer=self.execute(job)
        self.assertEqual(answer['status'],'done',answer['meta'].get('error'))
        context=json.dumps(self.calls,ensure_ascii=False);self.assertIn('1234',context);self.assertNotIn('999999',context)
        self.assertEqual(answer['meta']['rag_sources'][0]['source_id'],'chosen')

    def test_multiple_vision_batches_reuse_current_llm(self):
        from PIL import Image
        self.model['vision']={'enabled':True};self.model['max_refs']=2
        media=[]
        for i in range(3):
            stream=io.BytesIO();Image.new('RGB',(100,100),(i*80,10,20)).save(stream,'PNG')
            media.append(self.app.upload({'name':f'image-{i}.png','data':base64.b64encode(stream.getvalue()).decode()}))
        job=self.enqueue(media=media);self.calls=[];self.job=job;self.progress=0;self.fail=False;self.documents=False
        captions=[]
        def complete(messages,settings,cancel,**kw):
            if 'descriptions' in kw['schema']['properties']:
                images=[part for m in messages if isinstance(m['content'],list) for part in m['content'] if part['type']=='image_url'];captions.append(len(images))
                return json.dumps({'descriptions':['Figura verificata' for _ in images]}),'stop'
            return self.completion(messages,settings,cancel,**kw)
        with patch.object(self.app.engine,'require_model',return_value=self.model),patch.object(self.app.engine,'prepare'),patch.object(self.app.engine,'start_llama') as start,patch.object(self.app.engine,'completion',side_effect=complete):self.app.execute_job(job,self.cancel)
        answer=self.app.store.messages(job['chat_id'])[-1];self.assertEqual(answer['status'],'done',answer['meta'].get('error'));self.assertEqual(start.call_count,1);self.assertEqual(captions,[2,1])
        self.assertIn('Figura verificata',json.dumps(self.calls,ensure_ascii=False))

    def test_html_is_portable_and_sanitized(self):
        result=export_html(ROOT,self.app.data,{'title':'Slides','slide_format':'4:3','html':'<article class="h3-slide-page"><h1>Test</h1><script>alert(1)</script><img src="https://private.invalid/image"><p onclick="alert(1)">Text</p></article>'})
        source=(self.app.data/result['url'].lstrip('/')).read_text(encoding='utf-8')
        self.assertIn('1280px 960px',source);self.assertNotIn('<script',source);self.assertNotIn('onclick',source);self.assertNotIn('private.invalid',source);self.assertIn('script-src \'none\'',source)

if __name__=='__main__':unittest.main()
